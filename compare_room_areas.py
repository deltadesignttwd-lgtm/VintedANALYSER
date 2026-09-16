#!/usr/bin/env python3
"""Compare room areas between a "now" CSV and a "contract" CSV.

Workflow:
    1. From each Area cell (e.g. "49.13 m²"), take the first 4 characters
       from the left (e.g. "49.1", or "115." for three-digit areas).
    2. Convert that 4-character snippet into a numeric value (e.g. 49.1,
       115.0).
    3. Match rows between the two files by their room Number, compare the
       converted values, and write an Excel report with a Difference
       column, highlighting any row where the values don't match.

Usage:
    python compare_room_areas.py now.csv contract.csv --output comparison.xlsx

    python compare_room_areas.py sample_data/now_areas.csv \\
        sample_data/contract_areas.csv --output comparison.xlsx
"""

import argparse
import sys
from typing import Optional

import pandas as pd
from openpyxl.styles import PatternFill

MISMATCH_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
MISSING_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def extract_area_value(area: object) -> Optional[float]:
    """Step 1 + 2: take the first 4 characters from the left of the area
    string and convert them into a numeric value.

    "49.13 m²"  -> "49.1"  -> 49.1
    "115.20 m²" -> "115."  -> 115.0
    """
    if pd.isna(area):
        return None
    snippet = str(area).strip()[:4].rstrip(".")
    try:
        return float(snippet)
    except ValueError:
        return None


def load_areas(path: str, key_column: str, area_column: str) -> pd.DataFrame:
    df = pd.read_csv(path, dtype=str)
    if key_column not in df.columns or area_column not in df.columns:
        raise SystemExit(
            f"'{path}' must contain columns '{key_column}' and '{area_column}'. "
            f"Found: {list(df.columns)}"
        )
    df["_value"] = df[area_column].apply(extract_area_value)
    return df[[key_column, area_column, "_value"]]


def compare(
    now_path: str,
    contract_path: str,
    key_column: str,
    area_column: str,
) -> pd.DataFrame:
    now_df = load_areas(now_path, key_column, area_column).rename(
        columns={area_column: "Now Area", "_value": "Now Value"}
    )
    contract_df = load_areas(contract_path, key_column, area_column).rename(
        columns={area_column: "Contract Area", "_value": "Contract Value"}
    )

    merged = now_df.merge(contract_df, on=key_column, how="outer", indicator=True)

    def status(row: pd.Series) -> str:
        if row["_merge"] == "left_only":
            return "Missing in contract"
        if row["_merge"] == "right_only":
            return "Missing in now"
        if pd.isna(row["Now Value"]) or pd.isna(row["Contract Value"]):
            return "Unparseable area"
        if round(row["Now Value"] - row["Contract Value"], 6) != 0:
            return "Mismatch"
        return "Match"

    merged["Difference"] = merged["Now Value"] - merged["Contract Value"]
    merged["Status"] = merged.apply(status, axis=1)
    merged = merged.drop(columns=["_merge"])

    return merged[
        [key_column, "Now Area", "Contract Area", "Now Value", "Contract Value", "Difference", "Status"]
    ]


def write_report(df: pd.DataFrame, output_path: str) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Comparison")
        worksheet = writer.sheets["Comparison"]
        status_col_idx = df.columns.get_loc("Status") + 1  # 1-based for openpyxl

        for row_idx in range(2, worksheet.max_row + 1):
            status_value = worksheet.cell(row=row_idx, column=status_col_idx).value
            if status_value == "Mismatch":
                fill = MISMATCH_FILL
            elif status_value in ("Missing in now", "Missing in contract", "Unparseable area"):
                fill = MISSING_FILL
            else:
                continue
            for col_idx in range(1, worksheet.max_column + 1):
                worksheet.cell(row=row_idx, column=col_idx).fill = fill

        for col_idx, column in enumerate(df.columns, start=1):
            width = max(len(str(column)), df[column].astype(str).map(len).max()) + 2
            worksheet.column_dimensions[worksheet.cell(row=1, column=col_idx).column_letter].width = width


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare room areas between two CSVs.")
    parser.add_argument("now_csv", help="CSV with current room areas")
    parser.add_argument("contract_csv", help="CSV with contract room areas")
    parser.add_argument("--key", default="Number", help="Column used to match rows (default: Number)")
    parser.add_argument("--area-column", default="Area", help="Column holding the area value (default: Area)")
    parser.add_argument("--output", default="comparison.xlsx", help="Output .xlsx path (default: comparison.xlsx)")
    args = parser.parse_args()

    result = compare(args.now_csv, args.contract_csv, args.key, args.area_column)
    write_report(result, args.output)

    mismatches = (result["Status"] != "Match").sum()
    print(f"Compared {len(result)} rooms. {mismatches} row(s) flagged.")
    print(f"Report written to {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
